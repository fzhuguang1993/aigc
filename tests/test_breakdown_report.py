"""
tests/test_breakdown_report.py —— 拆解结果导出 Word（.docx）报告

拆解交付已从「回写任务表 12 列」改为「视频同目录同名 .docx」。这里锁三件事：
① 文档主名取视频文件名、落到视频所在目录（out_dir 可覆盖）；
② 正文各区块（标题/整体分析/逐字稿/分镜表/提示词）确实写进去、表格行数对得上；
③ 空结果也不能抛，任何残缺都能出一份合法 docx。
没装 python-docx 就整体跳过（它是拆解的可选重依赖）。
"""
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx import Document  # noqa: E402

from video_text_tools.breakdown import report  # noqa: E402
from video_text_tools.breakdown.models import (  # noqa: E402
    BreakdownResult, FrameAnalysis, SegmentPrompts, OverallAnalysis,
    TranscriptSegment)


def _result(tmp_path):
    """造一条尽量完整的拆解结果，视频落在 tmp_path 下（文档应与之同名同目录）。"""
    video = tmp_path / "骨胶原-9.27-142_11.mp4"
    video.write_bytes(b"x")
    return BreakdownResult(
        link="https://v.douyin.com/x",
        title="骨胶原口播",
        video_path=str(video),
        duration=42.0,
        shot_count=2,
        cost={"vision_calls": 1, "tokens": 26008},
        frame_analyses=[
            FrameAnalysis(idx=1, ts=0.0, shot_size="近景", camera="手持",
                          composition="居中", transition="无",
                          on_screen_text="开场", emotion="紧迫"),
            FrameAnalysis(idx=2, ts=4.0, shot_size="特写", camera="推"),
        ],
        transcript=[TranscriptSegment(0.0, 3.0, "每天两条胶原蛋白", [])],
        transcript_text="每天两条胶原蛋白",
        segments=[SegmentPrompts(index=1, time_range="00:00-00:03",
                                 visual_prompt="画面V", copy_prompt="文案C",
                                 shoot_prompt="复刻S", summary="开场")],
        overall=OverallAnalysis(hook_desc="前3秒抛痛点", hook_score="8",
                                factors="反差；悬念", emotion_curve="起伏",
                                formula="痛点+方案", blueprint="分步复刻"))


def test_write_report_named_after_video(tmp_path):
    res = _result(tmp_path)
    path = report.write_report(res)
    p = Path(path)
    assert p.name == "骨胶原-9.27-142_11.docx"      # 与视频同名
    assert p.parent == tmp_path                      # 落视频所在目录
    assert p.is_file()

    doc = Document(path)
    text = "\n".join(x.text for x in doc.paragraphs)
    assert "骨胶原口播" in text                       # 标题
    assert "前3秒抛痛点" in text                       # 一、整体分析
    assert "每天两条胶原蛋白" in text                  # 二、口播逐字稿
    assert "复刻S" in text                            # 四、三类提示词
    assert doc.tables, "应有分镜画面表"
    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "序号"
    assert len(table.rows) == 3                      # 表头 + 2 个分镜


def test_write_report_out_dir_override(tmp_path):
    sub = tmp_path / "自定义导出"
    res = _result(tmp_path)
    path = report.write_report(res, out_dir=str(sub))
    assert path.startswith(str(sub))
    assert (sub / "骨胶原-9.27-142_11.docx").is_file()


def test_write_report_minimal_no_error(tmp_path):
    """只有标题、其余全空也不能炸——任何残缺都出一份合法 docx。"""
    res = BreakdownResult(title="仅标题")
    path = report.write_report(res, out_dir=str(tmp_path))
    assert path.endswith("仅标题.docx")
    doc = Document(path)
    assert any("仅标题" in x.text for x in doc.paragraphs)
    assert not doc.tables                            # 无分镜则不建表


def test_safe_stem_prefers_video_and_sanitizes():
    """主名优先视频文件名；标题里的非法字符统一换成下划线。"""
    res = BreakdownResult(title="", video_path=r"C:\d\骨胶原_11.mp4")
    assert report._safe_stem(res) == "骨胶原_11"
    res2 = BreakdownResult(title="坏:名/称*?", video_path="")
    assert report._safe_stem(res2) == "坏_名_称__"
